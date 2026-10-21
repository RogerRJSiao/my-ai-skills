#!/usr/bin/env python3
"""
check_batch.py — check that a grading batch folder is ready before grading.

Read-only: nothing is created, moved, overwritten or deleted (temp_* templates
included). Run it after init_batch.py and after the roster, cloud listing and
submissions have been put in place.

Checks <data root>/grading/<domain>/<課程代碼>_<梯次>/ (resolved the same way
as init_batch.py):
  1. 00_input/manifest_source/  <課程>_三科_原始檔案<梯次>[_yyyymmdd_HHMMSS].csv
                                has data rows (not just the header)
  2. 00_input/roster/           <課程>_三科_學員名冊<梯次>.xlsx has data rows
  3. 00_input/submissions/      has first-level folders (count shown)
  4. 01_template/               every file in assets/01_template is present and
                                has data; shows its version and the 修改日
                                stored inside the workbook
  5. 02_work/Excel|Word|PPT     all empty
  6. 03_results/                holds only the templates from assets/03_results
Only files named for this batch are read; temp_* templates are skipped.
Each result is printed as soon as it is known. Exit code 0 = all passed.

Usage:
    python check_batch.py --batch 202609
    python check_batch.py --root D:\\claude_data --batch 202609
"""
import re
import csv
import sys
import argparse
import datetime as dt
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import openpyxl                                         # noqa: E402
from init_batch import (                                # noqa: E402
    ASSETS, GRADING_DIR, course_code, default_root, domain,
)

WORK_SUBDIRS = ("Excel", "Word", "PPT")
TEMPLATE_PREFIX = "temp_"


class Report:
    def __init__(self):
        self.failed = 0

    def ok(self, msg):
        print(f"  [OK]   {msg}", flush=True)

    def fail(self, msg):
        self.failed += 1
        print(f"  [FAIL] {msg}", flush=True)

    def info(self, msg):
        print(f"         {msg}", flush=True)


def section(title):
    print(f"\n{title}", flush=True)


def pick_one(folder: Path, pattern: re.Pattern, label: str, rep: Report) -> Path | None:
    """The single batch file in folder matching pattern; temp_* and Excel lock
    files are ignored, other files are reported but never read."""
    if not folder.is_dir():
        rep.fail(f"找不到資料夾 {folder}")
        return None
    files = [p for p in folder.iterdir() if p.is_file()
             and not p.name.startswith((TEMPLATE_PREFIX, "~$"))]
    hits = sorted(p for p in files if pattern.fullmatch(p.name))
    for p in sorted(set(files) - set(hits)):
        rep.info(f"略過（檔名與本梯次不符）：{p.name}")
    if not hits:
        rep.fail(f"找不到{label}")
        return None
    if len(hits) > 1:
        rep.fail(f"{label}有 {len(hits)} 份，只能有一份：{', '.join(p.name for p in hits)}")
        return None
    return hits[0]


def check_listing(dest: Path, course: str, batch: str, rep: Report):
    section("1. 00_input/manifest_source（原始檔案清單）")
    pat = re.compile(rf"{re.escape(course)}_三科_原始檔案{re.escape(batch)}(_\d{{8}}_\d{{6}})?\.csv")
    f = pick_one(dest / "00_input" / "manifest_source", pat,
                 f"{course}_三科_原始檔案{batch}[_yyyymmdd_HHMMSS].csv", rep)
    if not f:
        return
    try:
        with f.open(encoding="utf-8-sig", newline="") as fh:
            rows = list(csv.reader(fh))
    except (OSError, UnicodeDecodeError) as e:
        rep.fail(f"{f.name} 讀取失敗：{e}")
        return
    data = [r for r in rows[1:] if any(c.strip() for c in r)]
    rep.info(f"讀取 {f.name}：資料 {len(data)} 列（不含表頭）")
    if data:
        rep.ok(f"{f.name} 有資料")
    else:
        rep.fail(f"{f.name} 只有表頭，沒有資料")


def check_roster(dest: Path, course: str, batch: str, rep: Report):
    section("2. 00_input/roster（學員名冊）")
    pat = re.compile(rf"{re.escape(course)}_三科_學員名冊{re.escape(batch)}\.xlsx")
    f = pick_one(dest / "00_input" / "roster", pat, f"{course}_三科_學員名冊{batch}.xlsx", rep)
    if not f:
        return
    try:
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        n = sum(1 for r in ws.iter_rows(min_row=2, values_only=True)
                if any(v not in (None, "") for v in r))
        wb.close()
    except Exception as e:  # a broken workbook is a failed check, not a crash
        rep.fail(f"{f.name} 讀取失敗：{type(e).__name__}: {e}")
        return
    rep.info(f"讀取 {f.name}：資料 {n} 列（不含表頭）")
    if n:
        rep.ok(f"{f.name} 有資料")
    else:
        rep.fail(f"{f.name} 只有表頭，沒有資料")


def check_submissions(dest: Path, rep: Report):
    section("3. 00_input/submissions（學員作業）")
    folder = dest / "00_input" / "submissions"
    if not folder.is_dir():
        rep.fail(f"找不到資料夾 {folder}")
        return
    n = sum(1 for p in folder.iterdir() if p.is_dir())
    rep.info(f"第一層資料夾：{n} 個")
    if n:
        rep.ok("submissions 有學員資料夾")
    else:
        rep.fail("submissions 沒有任何資料夾")


def has_data(path: Path) -> bool:
    if path.suffix.lower() == ".csv":
        return path.stat().st_size > 0
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        return any(v not in (None, "") for ws in wb.worksheets
                   for r in ws.iter_rows(values_only=True) for v in r)
    finally:
        wb.close()


def check_templates(dest: Path, rep: Report):
    section("4. 01_template（評閱依據）")
    expected = sorted(p.name for p in (ASSETS / "01_template").iterdir() if p.is_file())
    folder = dest / "01_template"
    rep.info(f"應有 {len(expected)} 份：{', '.join(expected)}")
    for name in expected:
        f = folder / name
        if not f.is_file():
            rep.fail(f"缺少 {name}")
            continue
        m = re.search(r"_(v\d+(?:\.\d+)*)\.[^.]+$", name)
        version = m.group(1) if m else "（檔名沒有版本號）"
        try:
            if not has_data(f):
                rep.fail(f"{name} 沒有任何資料")
                continue
            modified = None
            if f.suffix.lower() == ".xlsx":
                wb = openpyxl.load_workbook(f, read_only=True)
                modified = wb.properties.modified
                wb.close()
        except Exception as e:
            rep.fail(f"{name} 讀取失敗：{type(e).__name__}: {e}")
            continue
        # docProps stores UTC; show local time like Excel's own file info
        when = (modified.replace(tzinfo=dt.timezone.utc).astimezone().strftime("%Y-%m-%d %H:%M")
                if modified else "（檔案內沒有記錄）")
        rep.ok(f"{name}　版本 {version}　修改日 {when}")


def check_work(dest: Path, rep: Report):
    section("5. 02_work（評閱工作區，必須淨空）")
    for sub in WORK_SUBDIRS:
        folder = dest / "02_work" / sub
        if not folder.is_dir():
            rep.fail(f"找不到資料夾 02_work/{sub}")
            continue
        items = sorted(folder.iterdir())
        if items:
            rep.fail(f"02_work/{sub} 不是空的（{len(items)} 項）：{', '.join(p.name for p in items[:5])}"
                     + (" …" if len(items) > 5 else ""))
        else:
            rep.ok(f"02_work/{sub} 是空的")


def check_results(dest: Path, rep: Report):
    section("6. 03_results（只能有樣板）")
    expected = {p.name for p in (ASSETS / "03_results").iterdir() if p.is_file()}
    folder = dest / "03_results"
    if not folder.is_dir():
        rep.fail(f"找不到資料夾 {folder}")
        return
    actual = {p.name for p in folder.iterdir()}
    for name in sorted(expected - actual):
        rep.fail(f"缺少 {name}")
    for name in sorted(actual - expected):
        rep.fail(f"多出不該有的項目：{name}")
    if actual == expected:
        rep.ok(f"只有 {len(expected)} 份樣板：{', '.join(sorted(expected))}")


def main():
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    ap = argparse.ArgumentParser(description="Check a grading batch folder before grading (read-only).")
    ap.add_argument("--root", default=None,
                    help="資料根目錄，例如 D:\\claude_data；省略時同 init_batch.py（local_config.json → 內建預設值）")
    ap.add_argument("--batch", required=True, help="梯次(yyyymm)，例如 202609")
    ap.add_argument("--course", default=None, help="課程代碼，預設讀 skill.toml 的 course_code")
    args = ap.parse_args()

    root_str, root_source = (args.root, "--root") if args.root else default_root()
    course = args.course or course_code()
    dest = Path(root_str) / GRADING_DIR / domain() / f"{course}_{args.batch}"
    print(f"檢查：{dest}（資料根目錄來源：{root_source}）", flush=True)
    if not dest.is_dir():
        sys.exit(f"找不到批次資料夾：{dest}。請先執行 init_batch.py --batch {args.batch}")

    rep = Report()
    #--1. manifest_source/ 要有 S06_三科_原始檔案yyyymm.csv，不只有表頭；只讀與目錄 S06_yyyymm 命名一致的檔案，讀取後即時顯示資料列數
    check_listing(dest, course, args.batch, rep)
    #--2. roster/ 要有 S06_三科_學員名冊yyyymm.xlsx，不只有表頭；只讀與目錄 S06_yyyymm 命名一致的檔案，讀取後即時顯示資料列數
    check_roster(dest, course, args.batch, rep)
    #--3. submissions/ 下方要有資料夾和檔案，只讀第一層的資料夾數量，讀取後即時顯示
    check_submissions(dest, rep)
    #--4. 01_template/ 要有三份檔案，每份都要有資料，並顯示版本號、修改日
    check_templates(dest, rep)
    #--5. 02_work/ 下方三個資料夾（Excel／Word／PPT）必須全部淨空
    check_work(dest, rep)
    #--6. 03_results/ 只能有 temp_S06_三科_學員檔案與成績計算yyyymm_v1.xlsx、temp_S06_單科_評分結果yyyymm_v1.xlsx 兩份
    check_results(dest, rep)

    print(f"\n{'全部通過' if not rep.failed else f'有 {rep.failed} 項未通過'}", flush=True)
    sys.exit(1 if rep.failed else 0)


if __name__ == "__main__":
    main()
