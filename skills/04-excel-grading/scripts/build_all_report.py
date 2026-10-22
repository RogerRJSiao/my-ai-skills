#!/usr/bin/env python3
"""
build_all_report.py — build the batch's 三科 學員檔案與成績計算 workbook.

  1. copies 03_results/temp_<課程>_三科_學員檔案與成績計算yyyymm_v1.xlsx to
     03_results/<課程>_三科_學員檔案與成績計算<梯次>_<yyyymmdd_HHMMSS>.xlsx
  2. sheet「manifest」: one row per file in
     00_input/manifest_source/<課程>_三科_原始檔案<梯次>[_yyyymmdd_HHMMSS].csv
     sheet「score」: one row per student of
     00_input/roster/<課程>_三科_學員名冊<梯次>.xlsx, in roster order
     (學員id, 姓名, 工號, 單位(來自名冊))
  3. fills Excel原始 / Word原始 / PPT原始 with 分數加總 of each subject's
     03_results/<課程>_<Excel|Word|PPT>_評分結果<梯次>_<yyyymmdd_HHMMSS>.xlsx
     (sheet「評分結果表」, matched by 工號; the newest file per subject).
     A missing subject file leaves that column blank; a student the subject
     file does not score gets 原始 0 as 歸檔異常, and so does any student whose
     file of that subject has 歸檔異常 or 命名異常 (原始 and the capped column).

manifest columns (rules shared with the MCP server's manifest):
  歸檔異常  path is not <課程>/<單位-工號-姓名>/<檔名>, the folder is not
            單位-6碼工號-姓名, the file is not 單位-6碼工號-姓名.docx|xlsx|pptx,
            the folder's 單位-工號-姓名 is not a roster row (單位 ignores case),
            or the file name differs from its folder (same rules as copy_to_work.py)
  姓名/工號/單位(來自檔名)/副檔名  split literally from the file name
  命名異常  the file name's 6碼工號 + 姓名 is not one roster student
  SHA256(前8碼)  of the local copy in 00_input/submissions/ (cloud path after <課程>/)
  檔案大小(KB)   from the listing CSV
  重複異常  same 副檔名 + SHA + 檔案大小 as another file: group code A, B, C...
  雲端讀檔時刻  from the listing file name; 評閱讀檔時刻  now
score columns besides the roster and 原始 (merged from the MCP server's
all_report rules):
  Excel / Word / PPT  0 when the student's file of that subject has 歸檔異常 or
                      命名異常 or was not submitted; otherwise 原始, capped at 100
  備註                「<科目>: <notes>」joined by「；」: 歸檔異常（未繳交）, 歸檔異常,
                      命名異常, 重複異常(<群組>), 多個<副檔名>檔, 單科超過100分以100分計,
                      歸檔異常（評分結果檔沒有此學員）
  Files are assigned to a student by the folder's 工號 first, then the file
  name's 工號. 總分 / 是否合格 are left as in the template.
Only the template copy is written; inputs are only read.

The MCP server (kd_mcp.core.grading.manifest / all_report) imports this file
for the manifest and roster rules, so the helpers raise ValueError instead of
exiting and only main() needs init_batch.py.

Usage:
    python build_all_report.py --batch 202609
    python build_all_report.py --root D:\\claude_data --batch 202609
"""
import re
import csv
import sys
import shutil
import hashlib
import argparse
import datetime as dt
from pathlib import Path, PurePosixPath

import openpyxl
from openpyxl.utils import get_column_letter

SUBJECTS = ("Excel", "Word", "PPT")
ALLOWED_EXTS = ("docx", "xlsx", "pptx")
STEM = r"[A-Za-z0-9]+-\d{6}-[^-.\s]+"           # 單位-6碼工號-姓名
FOLDER_RE = re.compile(STEM)
FILE_RE = re.compile(rf"{STEM}\.(?:{'|'.join(ALLOWED_EXTS)})", re.IGNORECASE)
MANIFEST_HEADERS = ["檔案id", "雲端檔案路徑", "歸檔異常", "姓名", "工號", "單位(來自檔名)", "命名異常",
                    "副檔名", "SHA256(前8碼)", "檔案大小(KB)", "重複異常", "雲端讀檔時刻", "評閱讀檔時刻"]
SCORE_HEADERS = ["學員id", "姓名", "工號", "單位(來自名冊)", "Excel原始", "Word原始", "PPT原始",
                 "Excel", "Word", "PPT", "總分", "是否合格", "備註"]
ROSTER_HEADERS = ["id", "姓名", "工號", "單位"]
TIME_FMT = "%Y%m%d %H:%M:%S"


# ── inputs ────────────────────────────────────────────────────────────────

def pick_one(folder: Path, pattern: str, label: str) -> Path:
    """The single batch file in folder matching pattern (temp_ / ~$ ignored)."""
    hits = sorted(p for p in folder.iterdir() if p.is_file()
                  and not p.name.startswith(("temp_", "~$")) and re.fullmatch(pattern, p.name))
    if len(hits) != 1:
        raise ValueError(f"{folder} 應有一份{label}，找到 {len(hits)} 份：{', '.join(p.name for p in hits) or '無'}")
    return hits[0]


def listing_pattern(course: str, batch: str) -> str:
    return rf"{re.escape(course)}_三科_原始檔案{re.escape(batch)}(_\d{{8}}_\d{{6}})?\.csv"


def roster_pattern(course: str, batch: str) -> str:
    return rf"{re.escape(course)}_三科_學員名冊{re.escape(batch)}\.xlsx"


def template_name(course: str) -> str:
    return f"temp_{course}_三科_學員檔案與成績計算yyyymm_v1.xlsx"


def report_name(course: str, batch: str, now: dt.datetime | None = None) -> str:
    return f"{course}_三科_學員檔案與成績計算{batch}_{(now or dt.datetime.now()):%Y%m%d_%H%M%S}.xlsx"


def read_roster(path: Path) -> list[tuple]:
    """[(id, 姓名, 工號, 單位)] exactly as in the roster, in roster order."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = wb.worksheets[0].iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows, ())]
    if header[:4] != ROSTER_HEADERS:
        raise ValueError(f"{path.name} 的表頭應為 {ROSTER_HEADERS}，實際為 {header[:4]}")
    out = [tuple(r[:4]) for r in rows if r and any(v not in (None, "") for v in r[:4])]
    wb.close()
    return out


# ── manifest ──────────────────────────────────────────────────────────────

def split_filename(filename: str):
    """(單位, 工號, 姓名, 副檔名) split literally; missing parts are None."""
    stem, dot, ext = filename.rpartition(".")
    if not dot:
        stem, ext = filename, None
    parts = stem.split("-", 2)
    return (*parts, ext) if len(parts) == 3 else (None, None, None, ext)


def sha8(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()[:8]


def group_code(n: int) -> str:
    """0 -> A, 25 -> Z, 26 -> AA ..."""
    code, n = "", n + 1
    while n:
        n, r = divmod(n - 1, 26)
        code = chr(65 + r) + code
    return code


def norm_path(key: str) -> str:
    """單位-工號-姓名 with 單位 upper-cased, so p1820 and P1820 compare equal."""
    unit, _, rest = key.partition("-")
    return f"{unit.upper()}-{rest}"


def build_manifest(listing: Path, submissions: Path, course: str, roster: list[tuple]):
    """(sheet rows, owner 工號 per row, (檔名, local copy or None) per row).
    The owner is the folder's 工號 when the folder is 單位-6碼工號-姓名,
    otherwise the file name's 工號."""
    m = re.search(r"(\d{8})_(\d{2})(\d{2})(\d{2})", listing.name)
    cloud_at = f"{m.group(1)} {m.group(2)}:{m.group(3)}:{m.group(4)}" if m else None
    review_at = dt.datetime.now().strftime(TIME_FMT)
    roster_keys = {(str(r[2]).strip(), str(r[1]).strip()) for r in roster}
    # 單位-工號-姓名 of every roster student (單位 upper-cased), as copy_to_work.py checks
    roster_paths = {f"{str(r[3]).strip().upper()}-{str(r[2]).strip()}-{str(r[1]).strip()}" for r in roster}
    with listing.open(encoding="utf-8-sig", newline="") as f:
        records = list(csv.DictReader(f))
    for col in ("檔名", "檔案路徑", "檔案大小(KB)"):
        if records and col not in records[0]:
            raise ValueError(f"{listing.name} 缺少欄位「{col}」")

    rows, keys, owners, files = [], [], [], []
    for i, rec in enumerate(records, start=1):
        cloud_path = rec["檔案路徑"].strip()
        filename = (rec.get("檔名") or PurePosixPath(cloud_path).name).strip()
        parts = PurePosixPath(cloud_path).parts
        rel = parts[parts.index(course) + 1:] if course in parts else ()
        folder = rel[0] if len(rel) == 2 else None
        unit, emp_id, name, ext = split_filename(filename)
        archive = not (folder and FOLDER_RE.fullmatch(folder) and FILE_RE.fullmatch(filename)
                       and norm_path(folder) in roster_paths                       # folder = a roster row
                       and norm_path(filename.rsplit(".", 1)[0]) == norm_path(folder))  # file = folder
        naming = not (emp_id and name and re.fullmatch(r"\d{6}", emp_id) and (emp_id, name) in roster_keys)
        local = submissions.joinpath(*rel) if rel else None
        sha = sha8(local) if local and local.is_file() else None
        try:
            size = float(rec.get("檔案大小(KB)"))
        except (TypeError, ValueError):
            size = None
        rows.append([i, cloud_path, "歸檔異常" if archive else None, name, emp_id, unit,
                     "命名異常" if naming else None, ext, sha, size, None, cloud_at, review_at])
        keys.append(((ext or "").lower(), sha, size) if sha else None)
        owners.append(folder.split("-")[1] if folder and FOLDER_RE.fullmatch(folder) else emp_id)
        files.append((filename, local))

    groups: dict[tuple, list[int]] = {}
    for idx, k in enumerate(keys):
        if k:
            groups.setdefault(k, []).append(idx)
    dups = sorted((g for g in groups.values() if len(g) > 1), key=lambda g: g[0])
    for n, g in enumerate(dups):
        for idx in g:
            rows[idx][10] = group_code(n)
    return rows, owners, files


def write_manifest_sheet(ws, rows: list[list]):
    """Rows 2.. of sheet manifest, with 檔案大小(KB) as 0.0 and an autofilter."""
    for i, row in enumerate(rows, start=2):
        for c, v in enumerate(row, start=1):
            ws.cell(row=i, column=c, value=v)
        ws.cell(row=i, column=10).number_format = "0.0"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(MANIFEST_HEADERS))}{len(rows) + 1}"


# ── per-student notes (same rules as the MCP server's all_report) ─────────

EXT_BY_SUBJECT = {"Excel": "xlsx", "Word": "docx", "PPT": "pptx"}
SCORE_CAP = 100
CAP_NOTE = f"單科超過{SCORE_CAP}分以{SCORE_CAP}分計"


def pick_file(files: list[list]) -> list:
    """Several manifest rows of one subject for a student: prefer one without
    歸檔/命名異常, then the lowest 檔案id."""
    return sorted(files, key=lambda r: (bool(r[2] or r[6]), r[0]))[0]


def file_status(manifest: list[list], owners: list, emp_id: str, subject: str):
    """(zero, notes) for one student and subject, from the manifest files
    assigned to the student. zero=True for 歸檔異常 / 命名異常 / 未繳交."""
    ext = EXT_BY_SUBJECT[subject]
    files = [r for r, o in zip(manifest, owners) if o == emp_id and (r[7] or "").lower() == ext]
    if not files:
        return True, ["歸檔異常（未繳交）"]
    f = pick_file(files)
    notes = [f"多個{ext}檔，採用檔案id {f[0]}"] if len(files) > 1 else []
    notes += [n for n in (f[2], f[6]) if n]
    if f[10]:
        notes.append(f"重複異常({f[10]})")
    return bool(f[2] or f[6]), notes


# ── subject scores ────────────────────────────────────────────────────────

RANGE_RE = re.compile(r"(?:'?([^'!]+)'?!)?\$?([A-Z]+)\$?(\d+):\$?([A-Z]+)\$?(\d+)")
CELL_RE = re.compile(r"\$?([A-Z]+)\$?(\d+)")


def _criteria(arg: str, ws):
    m = CELL_RE.fullmatch(arg)
    return ws[f"{m.group(1)}{m.group(2)}"].value if m else arg.strip('"')


def _col_values(wb, ws_default, rng: str) -> list:
    m = RANGE_RE.fullmatch(rng)
    if not m:
        raise ValueError(rng)
    ws = wb[m.group(1)] if m.group(1) else ws_default
    return [row[0].value for row in ws[f"{m.group(2)}{m.group(3)}:{m.group(4)}{m.group(5)}"]]


def eval_sum_formula(formula: str, wb, ws) -> float | None:
    """Evaluate the =SUMIFS(...) / =SUMIF(...) 分數加總 formulas the 評分結果
    files use, for workbooks saved without cached values."""
    m = re.fullmatch(r"=\s*(SUMIFS|SUMIF)\((.*)\)", formula.strip(), re.IGNORECASE)
    if not m:
        return None
    args = [a.strip() for a in m.group(2).split(",")]
    if m.group(1).upper() == "SUMIFS":
        sum_rng, pairs = args[0], list(zip(args[1::2], args[2::2]))
    else:
        sum_rng = args[2] if len(args) > 2 else args[0]
        pairs = [(args[0], args[1])]
    sums = _col_values(wb, ws, sum_rng)
    conds = [(_col_values(wb, ws, r), str(_criteria(c, ws)).strip()) for r, c in pairs]
    total = 0.0
    for i, v in enumerate(sums):
        if all(str(vals[i]).strip() == want for vals, want in conds) and isinstance(v, (int, float)):
            total += v
    return total


def subject_scores(results: Path, course: str, batch: str, subject: str):
    """({工號: 分數加總}, file used or None, warnings)."""
    pat = rf"{re.escape(course)}_{subject}_評分結果{re.escape(batch)}_\d{{8}}_\d{{6}}\.xlsx"
    files = sorted(p for p in results.iterdir() if p.is_file() and re.fullmatch(pat, p.name))
    if not files:
        return {}, None, []
    path, warns = files[-1], []
    if len(files) > 1:
        warns.append(f"有 {len(files)} 份，採用最新的 {path.name}")
    wb_f = openpyxl.load_workbook(path)                    # formulas
    wb_v = openpyxl.load_workbook(path, data_only=True)    # cached values (if Excel saved it)
    if "評分結果表" not in wb_f.sheetnames:
        return {}, path, [f"{path.name} 沒有活頁「評分結果表」，此科留空白"]
    ws_f, ws_v = wb_f["評分結果表"], wb_v["評分結果表"]
    header = [str(c.value).strip() if c.value is not None else "" for c in ws_f[1]]
    try:
        c_id, c_score = header.index("工號") + 1, header.index("分數加總") + 1
    except ValueError:
        return {}, path, [f"{path.name}「評分結果表」缺少「工號」或「分數加總」欄，此科留空白"]
    out = {}
    for r in range(2, ws_f.max_row + 1):
        emp = ws_f.cell(row=r, column=c_id).value
        if emp in (None, ""):
            continue
        val = ws_v.cell(row=r, column=c_score).value
        raw = ws_f.cell(row=r, column=c_score).value
        if val is None and isinstance(raw, str) and raw.startswith("="):
            val = eval_sum_formula(raw, wb_f, ws_f)
            if val is None:
                warns.append(f"第 {r} 列的分數加總公式無法計算：{raw}")
        elif val is None and isinstance(raw, (int, float)):
            val = raw
        emp = str(emp).strip()
        if emp in out:
            warns.append(f"工號 {emp} 出現多次，採用第一筆")
            continue
        out[emp] = int(val) if isinstance(val, float) and val.is_integer() else val
    return out, path, warns


# ── score sheet ───────────────────────────────────────────────────────────

def fill_score_sheet(ws, roster: list[tuple], manifest: list[list], owners: list,
                     results: Path, course: str, batch: str):
    """Rows 2.. of sheet score, in roster order. Returns (per-subject info:
    {subject, file, found, warnings, missing}, 備註 per roster student)."""
    for i, r in enumerate(roster, start=2):
        for c, v in enumerate(r, start=1):
            ws.cell(row=i, column=c, value=v)
    remarks = [[] for _ in roster]
    infos = []
    for subject in SUBJECTS:
        col = SCORE_HEADERS.index(f"{subject}原始") + 1
        cap_col = SCORE_HEADERS.index(subject) + 1
        scores, path, warns = subject_scores(results, course, batch, subject)
        info = {"subject": subject, "file": path, "found": 0, "warnings": warns, "missing": []}
        infos.append(info)
        if path is None:
            continue
        for i, r in enumerate(roster, start=2):
            emp = str(r[2]).strip()
            v = scores.get(emp)
            zero, notes = file_status(manifest, owners, emp, subject)
            #--沒評到分（評分結果檔沒有此學員）：視為歸檔異常，原始也以 0 計
            if v is None:
                info["missing"].append(str(r[1]))
                v, zero = 0, True
                if not any(n.startswith("歸檔異常") for n in notes):
                    notes.append("歸檔異常（評分結果檔沒有此學員）")
            else:
                info["found"] += 1
            #--歸檔異常、命名異常（含未繳交）：原始與 Excel／Word／PPT 都以 0 計
            if zero:
                v = 0
            ws.cell(row=i, column=col, value=v)
            #--Excel／Word／PPT：異常為 0，其餘為原始分數（上限 100）
            if zero:
                capped = 0
            else:
                capped = min(v, SCORE_CAP)
                if v > SCORE_CAP:
                    notes.append(CAP_NOTE)
            ws.cell(row=i, column=cap_col, value=capped)
            if notes:
                remarks[i - 2].append(f"{subject}: {'、'.join(notes)}")
    note_col = SCORE_HEADERS.index("備註") + 1
    remarks = ["；".join(parts) or None for parts in remarks]
    for i, text in enumerate(remarks, start=2):
        ws.cell(row=i, column=note_col, value=text)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(SCORE_HEADERS))}{len(roster) + 1}"
    return infos, remarks


# ── main ──────────────────────────────────────────────────────────────────

def check_headers(ws, expected: list[str]):
    actual = [ws.cell(row=1, column=c).value for c in range(1, len(expected) + 1)]
    if actual != expected:
        raise ValueError(f"範本活頁「{ws.title}」表頭不符：應為 {expected}，實際為 {actual}")


def main():
    try:
        run()
    except ValueError as e:
        sys.exit(str(e))


def run():
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from init_batch import GRADING_DIR, course_code, default_root, domain

    #--設定CLI
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    ap = argparse.ArgumentParser(description="Build the 三科 學員檔案與成績計算 workbook of one batch.")
    ap.add_argument("--root", default=None,
                    help="資料根目錄，例如 D:\\claude_data；省略時同 init_batch.py（local_config.json → 內建預設值）")
    ap.add_argument("--batch", required=True, help="梯次(yyyymm)，例如 202609")
    ap.add_argument("--course", default=None, help="課程代碼，預設讀 skill.toml 的 course_code")
    args = ap.parse_args()

    #--找到梯次資料夾與輸入檔
    root_str, root_source = (args.root, "--root") if args.root else default_root()
    course, batch = args.course or course_code(), args.batch
    dest = Path(root_str) / GRADING_DIR / domain() / f"{course}_{batch}"
    print(f"梯次：{dest}（資料根目錄來源：{root_source}）")
    if not dest.is_dir():
        sys.exit(f"找不到批次資料夾：{dest}")
    results = dest / "03_results"
    template = results / template_name(course)
    if not template.is_file():
        sys.exit(f"找不到樣板 {template}")
    listing = pick_one(dest / "00_input" / "manifest_source", listing_pattern(course, batch), "原始檔案清單")
    roster_path = pick_one(dest / "00_input" / "roster", roster_pattern(course, batch), "學員名冊")
    roster = read_roster(roster_path)

    #--1. 複製樣板成本次的檔案
    out = results / report_name(course, batch)
    if out.exists():
        sys.exit(f"{out.name} 已存在，請稍後再執行")
    shutil.copy2(template, out)
    print(f"1. 複製樣板 -> {out.name}")

    #--2. manifest（原始檔案清單）與 score 的學員資料（學員名冊）
    wb = openpyxl.load_workbook(out)
    ws_m, ws_s = wb["manifest"], wb["score"]
    check_headers(ws_m, MANIFEST_HEADERS)
    check_headers(ws_s, SCORE_HEADERS)
    manifest, owners, _ = build_manifest(listing, dest / "00_input" / "submissions", course, roster)
    write_manifest_sheet(ws_m, manifest)
    print(f"2. manifest：{listing.name} {len(manifest)} 個檔案（歸檔異常 "
          f"{sum(bool(r[2]) for r in manifest)}、命名異常 {sum(bool(r[6]) for r in manifest)}、"
          f"重複異常 {sum(bool(r[10]) for r in manifest)}）")
    print(f"   score：{roster_path.name} {len(roster)} 位學員")

    #--3. 各科 分數加總 -> Excel原始 / Word原始 / PPT原始
    print("3. 單科成績（評分結果表的分數加總）")
    infos, _ = fill_score_sheet(ws_s, roster, manifest, owners, results, course, batch)
    for info in infos:
        subject = info["subject"]
        if info["file"] is None:
            print(f"   {subject}：找不到 {course}_{subject}_評分結果{batch}_<yyyymmdd_HHMMSS>.xlsx，{subject}原始／{subject} 留空白")
            continue
        print(f"   {subject}：{info['file'].name}，填入 {info['found']}/{len(roster)} 位")
        for w in info["warnings"]:
            print(f"     注意：{w}")
        if info["missing"]:
            print(f"     名冊有、成績檔沒有（以 0 分計）：{'、'.join(info['missing'])}")

    wb.save(out)
    print(f"完成：{out}")


if __name__ == "__main__":
    main()
