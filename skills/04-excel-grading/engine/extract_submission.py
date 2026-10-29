"""Extract objective grading facts from one S06 Excel實作 submission.

The single implementation of the fact extraction: the xlsx_s06 engine imports
it, and it also runs on its own (only openpyxl needed):

    python engine/extract_submission.py <submission.xlsx> [--question <question.xlsx>]

prints the facts as JSON. Notable keys:

    chart_*:  chart_count, chart_types, bar_dir, is_3d, radar_style, has_data_table
    成績:     sort_conditions, filter_conditions, a1_merged_range, row_avg_uses_subtotal,
              g_values_descending, visible_data_rows, g16_cell
    pivot_薪資: host_sheet, location (the pivot may sit on a sheet other than 薪資;
              its source must be 薪資!A1:J13)
    pivot_薪資_rejected: pivots whose source failed that check, with the reason
    成績 (2.2): a1_horizontal, a1_merge_and_center
    top level: sheet_names, ambiguous_sheets (sheet names may carry a prefix/suffix),
               chartex_count (Excel 2016+ charts such as 直方圖, not inspected)

This module does NOT compute scores; see score_submission.py.
"""
from __future__ import annotations

import io
import posixpath
import re
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

import openpyxl
from openpyxl.utils import range_boundaries

NS = {
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
}
C = f"{{{NS['c']}}}"
M = f"{{{NS['main']}}}"
R = f"{{{NS['r']}}}"

SHEET_SALARY = "薪資"
SHEET_GRADES = "成績"
SHEET_VEHICLE = "車輛銷售"
SHEET_GENERAL = "武將能力"
EXPECTED_SHEETS = (SHEET_SALARY, SHEET_GRADES, SHEET_VEHICLE, SHEET_GENERAL)

# Raw-data cells compared against the original question file:
# column A (A2:A7) and row 2 (A2:E2), rubric 3.1 / 4.1.
_RAW_CELLS = (["A2", "A3", "A4", "A5", "A6", "A7"], ["B2", "C2", "D2", "E2"])
RAW_DATA_CELLS = {SHEET_VEHICLE: _RAW_CELLS, SHEET_GENERAL: _RAW_CELLS}

# The 薪資 data a qualifying pivot must use (rubric 1.1).
SALARY_SOURCE_RANGE = "A1:J13"


# ── workbook loading ─────────────────────────────────────────────────────

# Attributes Excel accepts but openpyxl's strict validation rejects, e.g.
# <a:latin pitchFamily="66"> in chart/drawing fonts (openpyxl allows <= 52).
_LENIENT_FIXES = (re.compile(rb'\s+pitchFamily="\d+"'),)


def load_workbook(path, data_only: bool = False):
    """openpyxl.load_workbook, retried once on a sanitised in-memory copy
    when openpyxl rejects drawing/chart XML that Excel itself opens fine.
    Only cosmetic attributes are removed; cell data is untouched, and the
    chart facts are read from the original zip by chart_info()."""
    try:
        return openpyxl.load_workbook(path, data_only=data_only)
    except ValueError:
        buf = io.BytesIO()
        with zipfile.ZipFile(path) as zin, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename.startswith(("xl/drawings/", "xl/charts/")) and item.filename.endswith(".xml"):
                    for rx in _LENIENT_FIXES:
                        data = rx.sub(b"", data)
                zout.writestr(item, data)
        buf.seek(0)
        return openpyxl.load_workbook(buf, data_only=data_only)


# ── OOXML package helpers ──────────────────────────────────────────────────

def _resolve_target(base_part: str, target: str) -> str:
    """Resolve a relationship Target relative to the part that owns it.
    Always uses '/' separators (zip member names), never os.sep."""
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(posixpath.join(posixpath.dirname(base_part), target))


def _rels_path(part: str) -> str:
    return posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels")


def _read_rels(zf: zipfile.ZipFile, part: str) -> list[dict]:
    try:
        root = ET.fromstring(zf.read(_rels_path(part)))
    except KeyError:
        return []
    return [
        {
            "Id": rel.attrib.get("Id"),
            "Type": rel.attrib.get("Type", ""),
            "Target": _resolve_target(part, rel.attrib.get("Target", "")),
        }
        for rel in root
    ]


def _workbook_rid_targets(zf: zipfile.ZipFile) -> dict[str, str]:
    return {r["Id"]: r["Target"] for r in _read_rels(zf, "xl/workbook.xml")}


def sheet_parts(zf: zipfile.ZipFile) -> dict[str, str]:
    """Map worksheet display name -> zip part path (e.g. 'xl/worksheets/sheet1.xml')."""
    wb_root = ET.fromstring(zf.read("xl/workbook.xml"))
    rid_to_target = _workbook_rid_targets(zf)
    out = {}
    for sheet in wb_root.iter(f"{M}sheet"):
        target = rid_to_target.get(sheet.attrib.get(f"{R}id"))
        if target:
            out[sheet.attrib.get("name")] = target
    return out


def resolve_sheet_names(names) -> tuple[dict[str, str], dict[str, list[str]]]:
    """Map each expected sheet to the student's actual sheet name.

    A sheet name may carry a prefix/suffix around the expected text
    (e.g. 「1.薪資」「成績_作答」). An exact match wins; otherwise the one
    sheet whose name contains the text is used. Several such sheets are
    ambiguous: none is picked, and they are returned for the caller to confirm."""
    names = [n for n in names if n]
    resolved, ambiguous = {}, {}
    for expected in EXPECTED_SHEETS:
        exact = [n for n in names if n.strip() == expected]
        found = exact or [n for n in names if expected in n]
        if len(found) == 1:
            resolved[expected] = found[0]
        elif found:
            ambiguous[expected] = found
    return resolved, ambiguous


# ── 考題1 薪資：pivot table ────────────────────────────────────────────────

def _pivot_targets(zf: zipfile.ZipFile, sheet_part: str) -> list[str]:
    return [r["Target"] for r in _read_rels(zf, sheet_part) if r["Type"].endswith("/pivotTable")]


def _table_sheets(zf: zipfile.ZipFile, parts: dict[str, str]) -> dict[str, tuple[str, str]]:
    """Map Excel table name (and displayName) -> (sheet it is defined on, its ref)."""
    out = {}
    for sheet, part in parts.items():
        for r in _read_rels(zf, part):
            if r["Type"].endswith("/table"):
                root = ET.fromstring(zf.read(r["Target"]))
                for key in ("name", "displayName"):
                    if root.attrib.get(key):
                        out[root.attrib[key]] = (sheet, root.attrib.get("ref"))
    return out


_SHEET_RANGE_RE = re.compile(r"'?(.+?)'?!\$?([A-Z]+)\$?(\d+):\$?([A-Z]+)\$?(\d+)")


def _defined_names(zf: zipfile.ZipFile) -> dict[str, tuple[str, str]]:
    """Map workbook-level defined name -> (sheet, range without $)."""
    out = {}
    for dn in ET.fromstring(zf.read("xl/workbook.xml")).iter(f"{M}definedName"):
        m = _SHEET_RANGE_RE.fullmatch((dn.text or "").strip())
        if m and dn.attrib.get("localSheetId") is None:
            out[dn.attrib.get("name")] = (m.group(1), f"{m.group(2)}{m.group(3)}:{m.group(4)}{m.group(5)}")
    return out


def _pivot_source(info: dict, tables: dict, names: dict) -> tuple[str | None, str | None, bool]:
    """(sheet, range without $, is an Excel table) of a pivot's source; (None, None, False)
    when it cannot be resolved."""
    src = info["source_name_or_range"]
    if info["_source_sheet"]:
        return info["_source_sheet"], (src or "").replace("$", "") or None, False
    if src in tables:
        return (*tables[src], True)
    if src in names:
        return (*names[src], False)
    return None, None, False


def find_salary_pivot(zf: zipfile.ZipFile, parts: dict[str, str],
                      salary_sheet: str | None) -> tuple[dict | None, list[dict]]:
    """(facts for the student's 薪資 pivot table or None, rejected pivots).

    The pivot may sit on 【薪資】 or on another sheet (Excel puts a new pivot
    on a new sheet by default), but its source must be 【薪資】!A1:J13. The
    sample on 【考題】 is skipped: it is hosted there and its cache reads an
    external workbook, so it never counts (rubric 1.1). Pivots that are not
    external but fail the source check are returned as `rejected`
    ({host_sheet, source_sheet, source_range, reason}) for the 評閱結果說明."""
    if salary_sheet is None:
        return None, []
    tables, names = _table_sheets(zf, parts), _defined_names(zf)
    hosts = [salary_sheet] + [s for s in parts if s != salary_sheet and "考題" not in s]
    rejected = []
    for host in hosts:
        for target in _pivot_targets(zf, parts[host]):
            info = pivot_info(zf, target)
            if info.pop("_external"):
                continue
            sheet, rng, is_table = _pivot_source(info, tables, names)
            info.pop("_source_sheet")
            if sheet is None:
                reason = "unresolved"
            elif sheet != salary_sheet:
                reason = "other_sheet"
            elif rng != SALARY_SOURCE_RANGE:
                reason = "wrong_range"
            else:
                return {**info, "source_is_named_table": is_table, "host_sheet": host}, rejected
            rejected.append({"host_sheet": host, "source_sheet": sheet, "source_range": rng,
                             "source_name_or_range": info["source_name_or_range"], "reason": reason})
    return None, rejected


def pivot_info(zf: zipfile.ZipFile, pivot_target: str) -> dict:
    """Facts for one pivotTable part."""
    pt_root = ET.fromstring(zf.read(pivot_target))
    cache_id = pt_root.attrib.get("cacheId")
    location = pt_root.find(f"{M}location")

    wb_root = ET.fromstring(zf.read("xl/workbook.xml"))
    rid_to_target = _workbook_rid_targets(zf)
    cache_target = None
    for pc in wb_root.iter(f"{M}pivotCache"):
        if pc.attrib.get("cacheId") == cache_id:
            cache_target = rid_to_target.get(pc.attrib.get(f"{R}id"))
            break

    cache_source_name = source_sheet = None
    external = False
    field_shared_items: dict[int, list] = {}
    if cache_target:
        cache_root = ET.fromstring(zf.read(cache_target))
        src = cache_root.find(f"{M}cacheSource/{M}worksheetSource")
        if src is not None:
            cache_source_name = src.attrib.get("name") or src.attrib.get("ref")
            source_sheet = src.attrib.get("sheet")
            external = src.attrib.get(f"{R}id") is not None  # source in another workbook
        for i, cf in enumerate(cache_root.findall(f"{M}cacheFields/{M}cacheField")):
            items = [s.attrib.get("v") for s in cf.findall(f"{M}sharedItems/{M}s")]
            if items:
                field_shared_items[i] = items

    pivot_fields = pt_root.findall(f"{M}pivotFields/{M}pivotField")
    row_field_idx = [f.attrib.get("x") for f in pt_root.findall(f"{M}rowFields/{M}field")]
    col_field_idx = [f.attrib.get("x") for f in pt_root.findall(f"{M}colFields/{M}field")]

    def visible_labels(field_index):
        if field_index is None:
            return None
        idx = int(field_index)
        if idx < 0 or idx >= len(pivot_fields):  # x="-2" is the Σ values field
            return None
        shared = field_shared_items.get(idx, [])
        labels = []
        for it in pivot_fields[idx].findall(f"{M}items/{M}item"):
            if it.attrib.get("t") == "default" or it.attrib.get("h") == "1":
                continue
            x = it.attrib.get("x") or "0"
            try:
                labels.append(shared[int(x)])
            except (IndexError, ValueError):
                labels.append(f"<item x={x}>")
        return labels

    data_fields = [
        {
            "name": df.attrib.get("name"),
            "subtotal": df.attrib.get("subtotal"),
            "numFmtId": df.attrib.get("numFmtId"),
        }
        for df in pt_root.findall(f"{M}dataFields/{M}dataField")
    ]

    return {
        "pivot_target": pivot_target,
        "source_name_or_range": cache_source_name,
        "row_field_index": row_field_idx,
        "col_field_index": col_field_idx,
        "row_visible_labels": visible_labels(row_field_idx[0]) if row_field_idx else None,
        "col_visible_labels": visible_labels(col_field_idx[0]) if col_field_idx else None,
        "data_fields": data_fields,
        "location": location.attrib.get("ref") if location is not None else None,
        "_source_sheet": source_sheet,
        "_external": external,
    }


def pivot_output_number_format(wb, sheet_name: str, location: str | None = None) -> list[str]:
    """Displayed number_format of numeric cells in the pivot output area
    (applyNumberFormats may be 0, so the cell's own style is what shows).
    `location` is the pivot's own range (e.g. "A3:C7"); without it J1:T10 is scanned."""
    ws = wb[sheet_name]
    min_col, min_row, max_col, max_row = range_boundaries(location) if location else (10, 1, 20, 10)
    formats = set()
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for c in row:
            if isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                formats.add(c.number_format)
    return sorted(formats)


# ── 考題3/4：charts ────────────────────────────────────────────────────────

def _chart_parts(zf: zipfile.ZipFile, sheet_part: str, rel_type: str = "/chart") -> list[str]:
    charts = []
    for d in _read_rels(zf, sheet_part):
        if d["Type"].endswith("/drawing"):
            charts += [r["Target"] for r in _read_rels(zf, d["Target"]) if r["Type"].endswith(rel_type)]
    return charts


def chartex_count(zf: zipfile.ZipFile, sheet_part: str) -> int:
    """Number of Excel 2016+ charts (chartEx: 直方圖, 瀑布圖…) on the sheet.
    Their XML differs from classic charts and is not inspected."""
    return len(_chart_parts(zf, sheet_part, "/chartEx"))


def _strcache_values(container):
    if container is None:
        return None
    cache = container.find(f"{C}strRef/{C}strCache")
    if cache is None:
        cache = container.find(f"{C}strLit")
    if cache is None:
        return None
    return [pt.find(f"{C}v").text for pt in cache.findall(f"{C}pt")]


def _single_chart_info(zf: zipfile.ZipFile, chart_target: str) -> dict:
    c_root = ET.fromstring(zf.read(chart_target))

    is_dynamic = c_root.find(f".//{C}numRef") is not None or c_root.find(f".//{C}strRef") is not None
    has_literal_only = (
        c_root.find(f".//{C}numLit") is not None or c_root.find(f".//{C}strLit") is not None
    ) and not is_dynamic

    # every chart group in the plot area, e.g. ["barChart", "lineChart"] for a combo chart
    plot_area = c_root.find(f".//{C}plotArea")
    chart_types = [el.tag.split("}")[-1] for el in (plot_area if plot_area is not None else [])
                   if el.tag.endswith("Chart")]
    # a 3-D clustered column chart counts as a 群組直條圖 (rubric 3.1)
    is_3d = "bar3DChart" in chart_types
    normalized = ["barChart" if t == "bar3DChart" else t for t in chart_types]
    chart_type = next((t for t in ("barChart", "radarChart") if t in normalized),
                      normalized[0] if normalized else None)

    first_cat = c_root.find(f".//{C}ser/{C}cat")
    categories = _strcache_values(first_cat) if first_cat is not None else None

    series_names = []
    for ser in c_root.findall(f".//{C}ser"):
        vals = _strcache_values(ser.find(f"{C}tx"))
        if vals:
            series_names.append(vals[0])

    cat_ref = None
    cat_strref = first_cat.find(f"{C}strRef/{C}f") if first_cat is not None else None
    if cat_strref is not None:
        cat_ref = cat_strref.text

    has_title = c_root.find(f".//{C}chart/{C}title") is not None
    auto_title_deleted = c_root.find(f".//{C}autoTitleDeleted")
    has_title = has_title and not (
        auto_title_deleted is not None and auto_title_deleted.attrib.get("val") == "1"
    )

    y_max = y_major_unit = gridlines_none = None
    val_ax = c_root.find(f".//{C}valAx")
    if val_ax is not None:
        scaling = val_ax.find(f"{C}scaling")
        if scaling is not None:
            mx = scaling.find(f"{C}max")
            y_max = mx.attrib.get("val") if mx is not None else None
        mu = val_ax.find(f"{C}majorUnit")
        y_major_unit = mu.attrib.get("val") if mu is not None else None
        gl = val_ax.find(f"{C}majorGridlines")
        gridlines_none = True if gl is None else gl.find(f".//{{{NS['a']}}}noFill") is not None

    has_data_labels = any(
        (sv := dl.find(f"{C}showVal")) is not None and sv.attrib.get("val") == "1"
        for dl in c_root.iter(f"{C}dLbls")
    )

    bar_dir = c_root.find(f".//{C}barDir")
    radar_style = c_root.find(f".//{C}radarStyle")

    return {
        "chart_target": chart_target,
        "chart_type": chart_type,
        "chart_types": chart_types,
        "is_dynamic": is_dynamic,
        "literal_only": has_literal_only,
        "category_ref": cat_ref,
        "categories": categories,
        "series_names": series_names,
        "y_max": y_max,
        "y_major_unit": y_major_unit,
        "gridlines_none": gridlines_none,
        "has_chart_title": has_title,
        "has_x_axis_title": c_root.find(f".//{C}catAx/{C}title") is not None,
        "has_y_axis_title": c_root.find(f".//{C}valAx/{C}title") is not None,
        "has_data_labels": has_data_labels,
        "marker_symbols": [m.attrib.get("val") for m in c_root.iter(f"{C}symbol")],
        # extras (not in the skill script)
        "bar_dir": bar_dir.attrib.get("val") if bar_dir is not None else None,
        "is_3d": is_3d,
        "radar_style": radar_style.attrib.get("val") if radar_style is not None else None,
        "has_data_table": c_root.find(f".//{C}plotArea/{C}dTable") is not None,
    }


def chart_info(zf: zipfile.ZipFile, sheet_part: str, preferred_type: str | None = None) -> dict | None:
    """Facts for the chart that should be graded on this sheet, or None.

    The rubric says: with several charts, only look at the dynamic chart of
    the requested type and ignore the rest. So pick the first chart that is
    dynamic and of `preferred_type`, falling back to the first chart."""
    charts = [_single_chart_info(zf, t) for t in _chart_parts(zf, sheet_part)]
    if not charts:
        return None
    chosen = next(
        (ch for ch in charts if ch["is_dynamic"] and ch["chart_type"] == preferred_type),
        charts[0],
    )
    return {**chosen, "chart_count": len(charts)}


# ── 考題2 成績 ─────────────────────────────────────────────────────────────

_ROUNDED_AVERAGE_RE = re.compile(r"=\s*(?:(?:ROUND|ROUNDUP|ROUNDDOWN)\(.*,\s*0\s*\)|INT\(.*\))\s*")


def shows_integer(cell) -> bool:
    """True if the cell's result is kept to whole numbers (rubric 2.3): an
    integer number format (0, #,##0, 0_) …) or ROUND(…, 0) / INT(…) around the formula."""
    section = (cell.number_format or "").split(";")[0]
    section = re.sub(r'"[^"]*"|\[[^\]]*\]|_.|\\.', "", section)   # literals, [colour], padding
    if re.search(r"[0#]", section) and not re.search(r"[.%Ee/]", section):
        return True
    return isinstance(cell.value, str) and bool(_ROUNDED_AVERAGE_RE.fullmatch(cell.value.upper()))


def grades_sheet_info(wb, xlsx_path: Path, sheet_name: str = SHEET_GRADES, wb_values=None) -> dict:
    """Facts for 考題2. `wb_values` is the workbook loaded with data_only=True;
    it supplies the cached 個人平均 values for the sort check (rubric 2.4)."""
    ws = wb[sheet_name]
    merged = [str(r) for r in ws.merged_cells.ranges]
    a1_merged = next((m for m in merged if m.startswith("A1:")), None)
    a1_val = ws["A1"].value
    a1_horizontal = ws["A1"].alignment.horizontal

    avg_row = next(
        (r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=1).value == "班平均"),
        None,
    )

    def is_average_formula(cell):
        return isinstance(cell.value, str) and "AVERAGE" in cell.value.upper()

    g_formulas_ok = all(is_average_formula(ws.cell(row=r, column=7)) for r in range(3, 16))

    row_avg_formulas_ok = row_avg_integer_fmt = None
    row_avg_uses_subtotal = False
    if avg_row:
        row_avg_formulas_ok = all(is_average_formula(ws.cell(row=avg_row, column=c)) for c in range(2, 7))
        row_avg_uses_subtotal = any(
            "SUBTOTAL(" in str(ws.cell(row=avg_row, column=c).value or "").upper() for c in range(2, 7)
        )
        row_avg_integer_fmt = all(shows_integer(ws.cell(row=avg_row, column=c)) for c in range(2, 7))
    # rubric 2.6 names the single cell G16, wherever 班平均 ends up
    g16 = ws["G16"]
    g16_cell = {"value": g16.value, "is_empty": g16.value in (None, "")}

    # rubric 2.5 names the fixed range A1:G16
    missing_border = []
    for r in range(1, 17):
        for col in range(1, 8):
            b = ws.cell(row=r, column=col).border
            if not (b.top.style or b.bottom.style or b.left.style or b.right.style):
                missing_border.append(ws.cell(row=r, column=col).coordinate)

    # openpyxl does not expose customFilters/sortState, read the table XML.
    sort_desc = filter_gt70 = None
    sort_conditions: list[dict] = []
    filter_conditions: list[dict] = []
    with zipfile.ZipFile(xlsx_path) as zf:
        for name in zf.namelist():
            if name.startswith("xl/tables/") and name.endswith(".xml"):
                root = ET.fromstring(zf.read(name))
                if root.attrib.get("name") == "表格2" or "個人平均" in ET.tostring(root, encoding="unicode"):
                    sort_desc = root.find(f".//{M}sortCondition") is not None
                    cf = root.find(f".//{M}customFilter")
                    filter_gt70 = cf is not None and cf.attrib.get("operator") == "greaterThan"
                    sort_conditions = [
                        {"ref": sc.attrib.get("ref"), "descending": sc.attrib.get("descending") == "1"}
                        for sc in root.iter(f"{M}sortCondition")
                    ]
                    filter_conditions = [
                        {
                            "col_id": fc.attrib.get("colId"),
                            "operator": f.attrib.get("operator", "equal"),
                            "val": f.attrib.get("val"),
                        }
                        for fc in root.iter(f"{M}filterColumn")
                        for f in fc.iter(f"{M}customFilter")
                    ]
                    # a ticked value list (<filters>) is not a 數字篩選 (rubric 2.4)
                    filter_conditions += [
                        {"col_id": fc.attrib.get("colId"), "operator": "values", "val": None}
                        for fc in root.iter(f"{M}filterColumn")
                        if fc.find(f"{M}filters") is not None
                    ]
    hidden_rows = [r for r in range(1, ws.max_row + 1) if ws.row_dimensions[r].hidden]
    # rows left after the filter: students + 班平均 (rubric 2.4 expects 8 + 1 = 9)
    visible_data_rows = sum(1 for r in range(3, (avg_row or 16) + 1) if not ws.row_dimensions[r].hidden)

    # rubric 2.4 checks the result: G3:G15 in descending order (cached values,
    # only present when Excel saved the file)
    g_values_descending = None
    if wb_values is not None:
        g_values = [wb_values[sheet_name].cell(row=r, column=7).value for r in range(3, 16)]
        if all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in g_values):
            g_values_descending = all(a >= b for a, b in zip(g_values, g_values[1:]))

    return {
        "freeze_panes": ws.freeze_panes,
        "merged_ranges": merged,
        "a1_value": a1_val,
        "a1_merged_range": a1_merged,
        "a1_horizontal": a1_horizontal,
        # 合併後置中 (rubric 2.2): A1:G1 merged and 水平置中 (跨欄置中 does not count)
        "a1_merge_and_center": a1_merged == "A1:G1" and a1_horizontal == "center",
        "header_text_preserved": a1_val == "經濟甲班 期中考成績表",
        "avg_row_index": avg_row,
        "g3_g15_average_formulas_ok": g_formulas_ok,
        "row_avg_formulas_ok": row_avg_formulas_ok,
        "row_avg_integer_format": row_avg_integer_fmt,
        "row_avg_uses_subtotal": row_avg_uses_subtotal,
        "g16_cell": g16_cell,
        "missing_border_cells": missing_border,
        "borders_complete": len(missing_border) == 0,
        "sorted_descending": sort_desc,
        "g_values_descending": g_values_descending,
        "filter_gt70_applied": filter_gt70,
        "hidden_rows": hidden_rows,
        "visible_data_rows": visible_data_rows,
        "sort_conditions": sort_conditions,
        "filter_conditions": filter_conditions,
    }


def raw_data_unchanged(wb_sub, wb_q, sheet_name: str, sub_sheet_name: str | None = None) -> dict:
    """Compare column A (A2:A7) and row 2 (A2:E2) with the question file.
    Leading/trailing whitespace is ignored: the question's D2 「程遠志」 carries a
    trailing Tab, and deleting it is not a change (rubric 4.1). A cell that is
    empty in the question must stay empty.
    `sub_sheet_name` is the student's sheet name when it has a prefix/suffix."""
    col_range, row2_range = RAW_DATA_CELLS[sheet_name]
    ws_sub, ws_q = wb_sub[sub_sheet_name or sheet_name], wb_q[sheet_name]

    def norm(v):
        v = v.strip() if isinstance(v, str) else v
        return None if v == "" else v

    diffs = []
    for coord in col_range + row2_range:
        v_sub, v_q = ws_sub[coord].value, ws_q[coord].value
        if norm(v_sub) != norm(v_q):
            diffs.append({"cell": coord, "submission": v_sub, "question": v_q})
    return {"unchanged": len(diffs) == 0, "diffs": diffs}


# ── entry point ────────────────────────────────────────────────────────────

def inspect_submission(xlsx_path: str | Path, question_path: str | Path | None = None) -> dict:
    """Return the facts as a dict (the JSON printed by the command line).

    Sheet names may have a prefix/suffix (see resolve_sheet_names);
    `sheet_names` maps each expected sheet to the name actually used.
    `missing_sheets` lists expected sheets that are absent and
    `ambiguous_sheets` those matched by several sheets; the caller should
    ask the user before grading such a file."""
    xlsx_path = Path(xlsx_path)
    out: dict = {"file": str(xlsx_path)}

    with zipfile.ZipFile(xlsx_path) as zf:
        parts = sheet_parts(zf)
        names, ambiguous = resolve_sheet_names(parts)
        out["sheet_names"] = names
        out["missing_sheets"] = [s for s in EXPECTED_SHEETS if s not in names and s not in ambiguous]
        out["ambiguous_sheets"] = ambiguous
        out["pivot_薪資"], out["pivot_薪資_rejected"] = find_salary_pivot(zf, parts, names.get(SHEET_SALARY))
        out["chart_車輛銷售"] = (
            chart_info(zf, parts[names[SHEET_VEHICLE]], "barChart") if SHEET_VEHICLE in names else None
        )
        out["chart_武將能力"] = (
            chart_info(zf, parts[names[SHEET_GENERAL]], "radarChart") if SHEET_GENERAL in names else None
        )
        out["chartex_count"] = {
            s: chartex_count(zf, parts[names[s]]) for s in (SHEET_VEHICLE, SHEET_GENERAL) if s in names
        }

    wb = load_workbook(xlsx_path, data_only=False)
    wb_sub = load_workbook(xlsx_path, data_only=True)
    if out["pivot_薪資"]:
        pivot = out["pivot_薪資"]
        pivot["value_number_formats"] = pivot_output_number_format(wb, pivot["host_sheet"], pivot["location"])
    out["成績"] = (
        grades_sheet_info(wb, xlsx_path, names[SHEET_GRADES], wb_sub) if SHEET_GRADES in names else None
    )

    if question_path:
        wb_q = load_workbook(question_path, data_only=True)
        for sheet in (SHEET_VEHICLE, SHEET_GENERAL):
            out[f"{sheet}_raw_data_check"] = (
                raw_data_unchanged(wb_sub, wb_q, sheet, names[sheet]) if sheet in names else None
            )
    return out


# ── command line ───────────────────────────────────────────────────────────

def main():
    import argparse
    import json
    import sys

    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    ap = argparse.ArgumentParser(description="Dump the facts needed to grade one S06 Excel submission (JSON).")
    ap.add_argument("submission")
    ap.add_argument("--question", default=None,
                    help="原始題目檔（如 assets/01_template/S06_excel_question_v1.xlsx），"
                         "用來檢查車輛銷售／武將能力的原始資料是否被更動")
    args = ap.parse_args()
    print(json.dumps(inspect_submission(args.submission, args.question),
                     ensure_ascii=False, indent=2, default=str))


if __name__ == "__main__":
    main()
