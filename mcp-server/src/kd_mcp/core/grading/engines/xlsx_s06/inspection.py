"""Extract objective grading facts from one S06 Excel實作 submission.

Ported from skills/04-excel-grading/scripts/inspect_submission.py so the MCP
server can import it as a library. Every key the script emits is kept with
the same meaning (tests/core/grading/engines/xlsx_s06/test_parity.py checks this), plus a few extra
keys excel_grading.scoring needs:

    chart_*:  chart_count, bar_dir, radar_style, has_data_table
    成績:     sort_conditions, filter_conditions

Like the script, this module does NOT compute scores; see scoring.py.
"""
from __future__ import annotations

import posixpath
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

import openpyxl

NS = {
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
}
C = f"{{{NS['c']}}}"
M = f"{{{NS['main']}}}"
R = f"{{{NS['r']}}}"

SHEET_SALARY = "薪資"
SHEET_GRADES = "成績"
SHEET_VEHICLE = "車輛銷售"
SHEET_GENERAL = "武將能力"
EXPECTED_SHEETS = (SHEET_SALARY, SHEET_GRADES, SHEET_VEHICLE, SHEET_GENERAL)

# Raw-data cells compared against the original question file.
RAW_DATA_CELLS = {
    SHEET_VEHICLE: (["A3", "A4", "A5", "A6"], ["B2", "C2", "D2"]),
    SHEET_GENERAL: (["A3", "A4", "A5", "A6", "A7"], ["B2", "C2", "D2", "E2"]),
}


# ── OOXML package helpers ──────────────────────────────────────────────────

def _resolve_target(base_part: str, target: str) -> str:
    """Resolve a relationship Target relative to the part that owns it.
    Always uses '/' separators (zip member names), never os.sep."""
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(posixpath.join(posixpath.dirname(base_part), target))


def _rels_path(part: str) -> str:
    return posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels")


def _read_rels(zf: zipfile.ZipFile, part: str) -> list[dict]:
    try:
        root = ET.fromstring(zf.read(_rels_path(part)))
    except KeyError:
        return []
    return [
        {
            "Id": rel.attrib.get("Id"),
            "Type": rel.attrib.get("Type", ""),
            "Target": _resolve_target(part, rel.attrib.get("Target", "")),
        }
        for rel in root
    ]


def _workbook_rid_targets(zf: zipfile.ZipFile) -> dict[str, str]:
    return {r["Id"]: r["Target"] for r in _read_rels(zf, "xl/workbook.xml")}


def sheet_parts(zf: zipfile.ZipFile) -> dict[str, str]:
    """Map worksheet display name -> zip part path (e.g. 'xl/worksheets/sheet1.xml')."""
    wb_root = ET.fromstring(zf.read("xl/workbook.xml"))
    rid_to_target = _workbook_rid_targets(zf)
    out = {}
    for sheet in wb_root.iter(f"{M}sheet"):
        target = rid_to_target.get(sheet.attrib.get(f"{R}id"))
        if target:
            out[sheet.attrib.get("name")] = target
    return out


# ── 考題1 薪資：pivot table ────────────────────────────────────────────────

def pivot_info(zf: zipfile.ZipFile, sheet_part: str) -> dict | None:
    """Facts for the pivotTable attached to this sheet's rels, or None.
    Pivot tables on other sheets (e.g. the sample on 【考題】) are ignored."""
    pivot_target = next(
        (r["Target"] for r in _read_rels(zf, sheet_part) if r["Type"].endswith("/pivotTable")),
        None,
    )
    if pivot_target is None:
        return None

    pt_root = ET.fromstring(zf.read(pivot_target))
    cache_id = pt_root.attrib.get("cacheId")

    wb_root = ET.fromstring(zf.read("xl/workbook.xml"))
    rid_to_target = _workbook_rid_targets(zf)
    cache_target = None
    for pc in wb_root.iter(f"{M}pivotCache"):
        if pc.attrib.get("cacheId") == cache_id:
            cache_target = rid_to_target.get(pc.attrib.get(f"{R}id"))
            break

    cache_source_name = None
    field_shared_items: dict[int, list] = {}
    if cache_target:
        cache_root = ET.fromstring(zf.read(cache_target))
        src = cache_root.find(f"{M}cacheSource/{M}worksheetSource")
        if src is not None:
            cache_source_name = src.attrib.get("name") or src.attrib.get("ref")
        for i, cf in enumerate(cache_root.findall(f"{M}cacheFields/{M}cacheField")):
            items = [s.attrib.get("v") for s in cf.findall(f"{M}sharedItems/{M}s")]
            if items:
                field_shared_items[i] = items

    pivot_fields = pt_root.findall(f"{M}pivotFields/{M}pivotField")
    row_field_idx = [f.attrib.get("x") for f in pt_root.findall(f"{M}rowFields/{M}field")]
    col_field_idx = [f.attrib.get("x") for f in pt_root.findall(f"{M}colFields/{M}field")]

    def visible_labels(field_index):
        if field_index is None:
            return None
        idx = int(field_index)
        if idx < 0 or idx >= len(pivot_fields):  # x="-2" is the Σ values field
            return None
        shared = field_shared_items.get(idx, [])
        labels = []
        for it in pivot_fields[idx].findall(f"{M}items/{M}item"):
            if it.attrib.get("t") == "default" or it.attrib.get("h") == "1":
                continue
            x = it.attrib.get("x") or "0"
            try:
                labels.append(shared[int(x)])
            except (IndexError, ValueError):
                labels.append(f"<item x={x}>")
        return labels

    data_fields = [
        {
            "name": df.attrib.get("name"),
            "subtotal": df.attrib.get("subtotal"),
            "numFmtId": df.attrib.get("numFmtId"),
        }
        for df in pt_root.findall(f"{M}dataFields/{M}dataField")
    ]

    return {
        "pivot_target": pivot_target,
        "source_name_or_range": cache_source_name,
        "source_is_named_table": bool(
            cache_source_name and not any(ch in cache_source_name for ch in "$:")
        ),
        "row_field_index": row_field_idx,
        "col_field_index": col_field_idx,
        "row_visible_labels": visible_labels(row_field_idx[0]) if row_field_idx else None,
        "col_visible_labels": visible_labels(col_field_idx[0]) if col_field_idx else None,
        "data_fields": data_fields,
    }


def pivot_output_number_format(wb, sheet_name: str) -> list[str]:
    """Displayed number_format of numeric cells in the pivot output area
    (applyNumberFormats may be 0, so the cell's own style is what shows)."""
    ws = wb[sheet_name]
    formats = set()
    for row in ws.iter_rows(min_row=1, max_row=10, min_col=10, max_col=20):
        for c in row:
            if isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                formats.add(c.number_format)
    return sorted(formats)


# ── 考題3/4：charts ────────────────────────────────────────────────────────

def _chart_parts(zf: zipfile.ZipFile, sheet_part: str) -> list[str]:
    charts = []
    for d in _read_rels(zf, sheet_part):
        if d["Type"].endswith("/drawing"):
            charts += [r["Target"] for r in _read_rels(zf, d["Target"]) if r["Type"].endswith("/chart")]
    return charts


def _strcache_values(container):
    if container is None:
        return None
    cache = container.find(f"{C}strRef/{C}strCache")
    if cache is None:
        cache = container.find(f"{C}strLit")
    if cache is None:
        return None
    return [pt.find(f"{C}v").text for pt in cache.findall(f"{C}pt")]


def _single_chart_info(zf: zipfile.ZipFile, chart_target: str) -> dict:
    c_root = ET.fromstring(zf.read(chart_target))

    is_dynamic = c_root.find(f".//{C}numRef") is not None or c_root.find(f".//{C}strRef") is not None
    has_literal_only = (
        c_root.find(f".//{C}numLit") is not None or c_root.find(f".//{C}strLit") is not None
    ) and not is_dynamic

    chart_type = None
    for tag in ("barChart", "radarChart", "lineChart", "pieChart"):
        if c_root.find(f".//{C}{tag}") is not None:
            chart_type = tag
            break

    first_cat = c_root.find(f".//{C}ser/{C}cat")
    categories = _strcache_values(first_cat) if first_cat is not None else None

    series_names = []
    for ser in c_root.findall(f".//{C}ser"):
        vals = _strcache_values(ser.find(f"{C}tx"))
        if vals:
            series_names.append(vals[0])

    cat_ref = None
    cat_strref = first_cat.find(f"{C}strRef/{C}f") if first_cat is not None else None
    if cat_strref is not None:
        cat_ref = cat_strref.text

    has_title = c_root.find(f".//{C}chart/{C}title") is not None
    auto_title_deleted = c_root.find(f".//{C}autoTitleDeleted")
    has_title = has_title and not (
        auto_title_deleted is not None and auto_title_deleted.attrib.get("val") == "1"
    )

    y_max = y_major_unit = gridlines_none = None
    val_ax = c_root.find(f".//{C}valAx")
    if val_ax is not None:
        scaling = val_ax.find(f"{C}scaling")
        if scaling is not None:
            mx = scaling.find(f"{C}max")
            y_max = mx.attrib.get("val") if mx is not None else None
        mu = val_ax.find(f"{C}majorUnit")
        y_major_unit = mu.attrib.get("val") if mu is not None else None
        gl = val_ax.find(f"{C}majorGridlines")
        gridlines_none = True if gl is None else gl.find(f".//{{{NS['a']}}}noFill") is not None

    has_data_labels = any(
        (sv := dl.find(f"{C}showVal")) is not None and sv.attrib.get("val") == "1"
        for dl in c_root.iter(f"{C}dLbls")
    )

    bar_dir = c_root.find(f".//{C}barDir")
    radar_style = c_root.find(f".//{C}radarStyle")

    return {
        "chart_target": chart_target,
        "chart_type": chart_type,
        "is_dynamic": is_dynamic,
        "literal_only": has_literal_only,
        "category_ref": cat_ref,
        "categories": categories,
        "series_names": series_names,
        "y_max": y_max,
        "y_major_unit": y_major_unit,
        "gridlines_none": gridlines_none,
        "has_chart_title": has_title,
        "has_x_axis_title": c_root.find(f".//{C}catAx/{C}title") is not None,
        "has_y_axis_title": c_root.find(f".//{C}valAx/{C}title") is not None,
        "has_data_labels": has_data_labels,
        "marker_symbols": [m.attrib.get("val") for m in c_root.iter(f"{C}symbol")],
        # extras (not in the skill script)
        "bar_dir": bar_dir.attrib.get("val") if bar_dir is not None else None,
        "radar_style": radar_style.attrib.get("val") if radar_style is not None else None,
        "has_data_table": c_root.find(f".//{C}plotArea/{C}dTable") is not None,
    }


def chart_info(zf: zipfile.ZipFile, sheet_part: str, preferred_type: str | None = None) -> dict | None:
    """Facts for the chart that should be graded on this sheet, or None.

    The rubric says: with several charts, only look at the dynamic chart of
    the requested type and ignore the rest. So pick the first chart that is
    dynamic and of `preferred_type`, falling back to the first chart."""
    charts = [_single_chart_info(zf, t) for t in _chart_parts(zf, sheet_part)]
    if not charts:
        return None
    chosen = next(
        (ch for ch in charts if ch["is_dynamic"] and ch["chart_type"] == preferred_type),
        charts[0],
    )
    return {**chosen, "chart_count": len(charts)}


# ── 考題2 成績 ─────────────────────────────────────────────────────────────

def grades_sheet_info(wb, xlsx_path: Path) -> dict:
    ws = wb[SHEET_GRADES]
    merged = [str(r) for r in ws.merged_cells.ranges]
    a1_val = ws["A1"].value
    centered_across = ws["A1"].alignment.horizontal in ("center", "centerContinuous")

    avg_row = next(
        (r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=1).value == "班平均"),
        None,
    )

    def is_average_formula(cell):
        return isinstance(cell.value, str) and "AVERAGE" in cell.value.upper()

    g_formulas_ok = all(is_average_formula(ws.cell(row=r, column=7)) for r in range(3, 16))

    row_avg_formulas_ok = row_avg_integer_fmt = g_avg_row_value = None
    if avg_row:
        row_avg_formulas_ok = all(is_average_formula(ws.cell(row=avg_row, column=c)) for c in range(2, 7))
        row_avg_integer_fmt = all(
            ws.cell(row=avg_row, column=c).number_format in ("0", "0_)") for c in range(2, 7)
        )
        g_cell = ws.cell(row=avg_row, column=7)
        g_avg_row_value = {"value": g_cell.value, "is_empty": g_cell.value in (None, "")}

    last_row = avg_row or 16
    missing_border = []
    for r in range(1, last_row + 1):
        for col in range(1, 8):
            b = ws.cell(row=r, column=col).border
            if not (b.top.style or b.bottom.style or b.left.style or b.right.style):
                missing_border.append(ws.cell(row=r, column=col).coordinate)

    # openpyxl does not expose customFilters/sortState, read the table XML.
    sort_desc = filter_gt70 = None
    sort_conditions: list[dict] = []
    filter_conditions: list[dict] = []
    with zipfile.ZipFile(xlsx_path) as zf:
        for name in zf.namelist():
            if name.startswith("xl/tables/") and name.endswith(".xml"):
                root = ET.fromstring(zf.read(name))
                if root.attrib.get("name") == "表格2" or "個人平均" in ET.tostring(root, encoding="unicode"):
                    sort_desc = root.find(f".//{M}sortCondition") is not None
                    cf = root.find(f".//{M}customFilter")
                    filter_gt70 = cf is not None and cf.attrib.get("operator") == "greaterThan"
                    sort_conditions = [
                        {"ref": sc.attrib.get("ref"), "descending": sc.attrib.get("descending") == "1"}
                        for sc in root.iter(f"{M}sortCondition")
                    ]
                    filter_conditions = [
                        {
                            "col_id": fc.attrib.get("colId"),
                            "operator": f.attrib.get("operator", "equal"),
                            "val": f.attrib.get("val"),
                        }
                        for fc in root.iter(f"{M}filterColumn")
                        for f in fc.iter(f"{M}customFilter")
                    ]
    hidden_rows = [r for r in range(1, ws.max_row + 1) if ws.row_dimensions[r].hidden]

    return {
        "freeze_panes": ws.freeze_panes,
        "merged_ranges": merged,
        "a1_value": a1_val,
        "a1_center_or_merged": bool(merged) or centered_across,
        "header_text_preserved": a1_val == "經濟甲班 期中考成績表",
        "avg_row_index": avg_row,
        "g3_g15_average_formulas_ok": g_formulas_ok,
        "row_avg_formulas_ok": row_avg_formulas_ok,
        "row_avg_integer_format": row_avg_integer_fmt,
        "g_avg_row_cell": g_avg_row_value,
        "missing_border_cells": missing_border,
        "borders_complete": len(missing_border) == 0,
        "sorted_descending": sort_desc,
        "filter_gt70_applied": filter_gt70,
        "hidden_rows": hidden_rows,
        "sort_conditions": sort_conditions,
        "filter_conditions": filter_conditions,
    }


def raw_data_unchanged(wb_sub, wb_q, sheet_name: str) -> dict:
    """Compare A-column categories and row-2 headers with the question file."""
    col_range, row2_range = RAW_DATA_CELLS[sheet_name]
    ws_sub, ws_q = wb_sub[sheet_name], wb_q[sheet_name]
    diffs = []
    for coord in col_range + row2_range:
        v_sub, v_q = ws_sub[coord].value, ws_q[coord].value
        if v_sub != v_q:
            diffs.append({"cell": coord, "submission": v_sub, "question": v_q})
    return {"unchanged": len(diffs) == 0, "diffs": diffs}


# ── entry point ────────────────────────────────────────────────────────────

def inspect_submission(xlsx_path: str | Path, question_path: str | Path | None = None) -> dict:
    """Return the same JSON structure as inspect_submission.py (plus extras).

    `missing_sheets` lists expected sheets that are absent; the caller should
    ask the user before grading such a file."""
    xlsx_path = Path(xlsx_path)
    out: dict = {"file": str(xlsx_path)}

    with zipfile.ZipFile(xlsx_path) as zf:
        parts = sheet_parts(zf)
        out["missing_sheets"] = [s for s in EXPECTED_SHEETS if s not in parts]
        out["pivot_薪資"] = pivot_info(zf, parts[SHEET_SALARY]) if SHEET_SALARY in parts else None
        out["chart_車輛銷售"] = (
            chart_info(zf, parts[SHEET_VEHICLE], "barChart") if SHEET_VEHICLE in parts else None
        )
        out["chart_武將能力"] = (
            chart_info(zf, parts[SHEET_GENERAL], "radarChart") if SHEET_GENERAL in parts else None
        )

    wb = openpyxl.load_workbook(xlsx_path, data_only=False)
    if out["pivot_薪資"]:
        out["pivot_薪資"]["value_number_formats"] = pivot_output_number_format(wb, SHEET_SALARY)
    out["成績"] = grades_sheet_info(wb, xlsx_path) if SHEET_GRADES in wb.sheetnames else None

    if question_path:
        wb_sub = openpyxl.load_workbook(xlsx_path, data_only=True)
        wb_q = openpyxl.load_workbook(question_path, data_only=True)
        for sheet in (SHEET_VEHICLE, SHEET_GENERAL):
            out[f"{sheet}_raw_data_check"] = (
                raw_data_unchanged(wb_sub, wb_q, sheet) if sheet in wb_sub.sheetnames else None
            )
    return out
